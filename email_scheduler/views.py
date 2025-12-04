# exam/views.py

import openpyxl
from django.http import HttpResponse
from django.db import connection


def export_tests_excel(request):
    """
    Export test summary to Excel based on date range.
    URL examples:
    /export-tests-excel?start_date=2025-09-21&end_date=2025-09-23
    /export-tests-excel?start_date=2025-09-23
    """

    # Get query params
    start_date = request.GET.get("start_date")  # e.g. 2025-09-21
    end_date = request.GET.get("end_date")      # e.g. 2025-09-23

    # Base SQL
    base_sql = """
        SELECT ui.date_of_exam,
               ui.start_time_slot,
               t.test_type,
               t.suspicious_activity_count,
               COUNT(t.suspicious_activity_count) AS total_suspicious_users,
               t.resolution_mismatch_count,
               COUNT(t.resolution_mismatch_count) AS total_resolution_mismatch_users
        FROM tests t
        INNER JOIN user_invites ui ON ui.user_id = t.user_id
        WHERE t.is_submitted = 1
    """

    params = []

    if start_date and end_date:
        base_sql += " AND ui.date_of_exam BETWEEN %s AND %s"
        params.extend([start_date, end_date])
    elif start_date:
        base_sql += " AND ui.date_of_exam = %s"
        params.append(start_date)
    elif end_date:
        base_sql += " AND ui.date_of_exam = %s"
        params.append(end_date)
    else:
        return HttpResponse("Please provide at least one date (start_date or end_date).", status=400)

    base_sql += """
        GROUP BY t.test_type, t.suspicious_activity_count, 
                 t.resolution_mismatch_count, ui.start_time_slot, ui.date_of_exam
        ORDER BY t.test_type ASC
    """

    # Run query
    with connection.cursor() as cursor:
        cursor.execute(base_sql, params)
        rows = cursor.fetchall()
        columns = [col[0] for col in cursor.description]

    # Build Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tests Summary"

    # Add headers
    ws.append(columns)

    # Add rows
    for row in rows:
        ws.append(row)

    # Response
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    if start_date and end_date:
        filename = f"tests_summary_{start_date}_to_{end_date}.xlsx"
    else:
        filename = f"tests_summary_{start_date or end_date}.xlsx"

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def export_exam_summary(request):
    """
    Export exam summary with counts:
    - Total Invited: Count from user_invites table, grouped by date_of_exam and start_time_slot
    - Exam Attended Count: Per test type based on start_time not null
    - Final Submission: Per test type (is_submitted = 1)
    - Auto Submission: Per test type (suspicious >= 9 OR resolution >= 3)

    Example:
    /export-exam-summary?start_date=2025-09-21&end_date=2025-09-23
    /export-exam-summary?start_date=2025-09-23
    """

    # Get query params
    start_date = request.GET.get("start_date")
    end_date = request.GET.get("end_date")

    # Validate date inputs
    if not start_date and not end_date:
        return HttpResponse("Please provide at least one date (start_date or end_date).", status=400)

    # Query to calculate total_invited per date_of_exam and start_time_slot
    invited_count_sql = """
        SELECT 
            date_of_exam,
            start_time_slot,
            COUNT(DISTINCT user_id) as total_invited
        FROM user_invites
        WHERE 1=1
    """
    
    invited_params = []
    if start_date and end_date:
        invited_count_sql += " AND date_of_exam BETWEEN %s AND %s"
        invited_params.extend([start_date, end_date])
    elif start_date:
        invited_count_sql += " AND date_of_exam = %s"
        invited_params.append(start_date)
    elif end_date:
        invited_count_sql += " AND date_of_exam = %s"
        invited_params.append(end_date)

    invited_count_sql += """
        GROUP BY date_of_exam, start_time_slot
    """

    # Execute query to get total_invited counts
    invited_counts = {}
    with connection.cursor() as cursor:
        cursor.execute(invited_count_sql, invited_params)
        for row in cursor.fetchall():
            date_of_exam, start_time_slot, total_invited = row
            invited_counts[(date_of_exam, start_time_slot)] = total_invited

    # Main SQL query
    base_sql = """
        SELECT 
            ui.date_of_exam,
            ui.start_time_slot,
            t.test_type,

            -- Exam attended count per test type
            COUNT(CASE WHEN t.start_time IS NOT NULL THEN t.user_id ELSE NULL END) AS exam_attended_count,

            -- Final submission per test type
            COUNT(CASE WHEN t.is_submitted = 1 THEN t.user_id ELSE NULL END) AS final_submission,

            -- Auto submission per test type
            COUNT(CASE 
                    WHEN t.suspicious_activity_count >= 9 
                     OR t.resolution_mismatch_count >= 3 
                    THEN t.user_id ELSE NULL 
                END) AS auto_submission

        FROM user_invites ui
        LEFT JOIN tests t ON ui.user_id = t.user_id
        WHERE 1=1
        AND t.test_type IS NOT NULL
    """

    params = []
    if start_date and end_date:
        base_sql += " AND ui.date_of_exam BETWEEN %s AND %s"
        params.extend([start_date, end_date])
    elif start_date:
        base_sql += " AND ui.date_of_exam = %s"
        params.append(start_date)
    elif end_date:
        base_sql += " AND ui.date_of_exam = %s"
        params.append(end_date)

    base_sql += """
        GROUP BY 
            ui.date_of_exam,
            ui.start_time_slot,
            t.test_type
        ORDER BY 
            ui.date_of_exam ASC,
            ui.start_time_slot ASC,
            t.test_type ASC
    """

    # Run main query
    with connection.cursor() as cursor:
        cursor.execute(base_sql, params)
        rows = cursor.fetchall()
        columns = [col[0] for col in cursor.description]

    # Insert total_invited into rows
    updated_rows = []
    total_invited_idx = columns.index('test_type')  # Insert total_invited before test_type
    columns.insert(total_invited_idx, 'total_invited')  # Update headers

    for row in rows:
        row = list(row)
        date_of_exam, start_time_slot = row[0], row[1]
        total_invited = invited_counts.get((date_of_exam, start_time_slot), 0)
        row.insert(total_invited_idx, total_invited)
        updated_rows.append(row)

    # Build Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exam Summary"

    ws.append(columns)  # headers
    for row in updated_rows:
        ws.append(row)

    # Response
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    if start_date and end_date:
        filename = f"exam_summary_{start_date}_to_{end_date}.xlsx"
    else:
        filename = f"exam_summary_{start_date or end_date}.xlsx"

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def export_browser_details(request):
    start_date = request.GET.get("start_date")
    end_date = request.GET.get("end_date")

    if not start_date and not end_date:
        return HttpResponse("Please provide at least start_date or end_date (?start_date=YYYY-MM-DD&end_date=YYYY-MM-DD).", status=400)

    base_sql = """
        SELECT 
            ui.date_of_exam, 
            ui.start_time_slot, 
            t.test_type, 
            us.registration_id, 
            us.email, 
            t.browser_used, 
            t.suspicious_activity_count,  
            t.resolution_mismatch_count 
        FROM users us 
        JOIN tests t ON us.user_id = t.user_id 
        JOIN user_invites ui ON t.user_id = ui.user_id  
        WHERE 1=1
          AND (t.test_type IN ('GT', 'TT'))
          AND (t.suspicious_activity_count >= 9 OR t.resolution_mismatch_count >= 3)
    """

    params = []

    if start_date and end_date:
        base_sql += " AND ui.date_of_exam BETWEEN %s AND %s"
        params.extend([start_date, end_date])
    elif start_date:  # only start_date given
        base_sql += " AND ui.date_of_exam >= %s"
        params.append(start_date)
    elif end_date:  # only end_date given
        base_sql += " AND ui.date_of_exam <= %s"
        params.append(end_date)

    base_sql += " ORDER BY ui.date_of_exam ASC, t.test_type ASC, ui.start_time_slot ASC"

    with connection.cursor() as cursor:
        cursor.execute(base_sql, params)
        rows = cursor.fetchall()
        columns = [col[0] for col in cursor.description]

    if not rows:
        return HttpResponse(f"No data found for given date range.", status=404)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Browser Details"
    ws.append(columns)

    for row in rows:
        ws.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    if start_date and end_date:
        filename = f"tests_browser_details_{start_date}_to_{end_date}.xlsx"
    elif start_date:
        filename = f"tests_browser_details_from_{start_date}.xlsx"
    else:
        filename = f"tests_browser_details_until_{end_date}.xlsx"

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

